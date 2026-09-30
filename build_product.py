"""Build a sellable funder list (ZIP with XLSX + CSV + README) from the grants CSV.

Usage:
  python build_product.py --grants data/grants_2026.csv --keyword literacy --keyword "adult reading" \
      --title "Literacy funders 2026" --slug literacy-funders-2026 --out dist

The list is a filtered view of what foundations filed on Form 990-PF, Part XV.
Nothing is added, guessed or scored; every row keeps its source object ID.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from pfgrants.query import filter_rows, read_csv, summarize_by_foundation, write_csv, GRANT_COLUMNS

FOUNDATION_COLS = [
    ("foundation_name", "Foundation", 42), ("foundation_ein", "EIN", 12),
    ("foundation_city", "City", 18), ("foundation_state", "State", 7),
    ("tax_year", "Tax year", 9), ("preselected_only", "Pre-selected only (Part XV line 2)", 14),
    ("application_contact", "Application contact (line 2a)", 30),
    ("application_address", "Application address (line 2a)", 36),
    ("application_deadlines", "Deadlines (line 2c)", 24), ("website", "Website", 26),
    ("fmv_assets_eoy", "Assets at year end (FMV, USD)", 16),
    ("matching_grants", "Matching grants", 10), ("matching_amount", "Matching amount (USD)", 16),
    ("recipients", "Matching recipients", 60), ("source_url", "Filing (ProPublica)", 40),
]
GRANT_COLS = [
    ("foundation_name", "Foundation", 40), ("foundation_state", "Fdn state", 8),
    ("grant_kind", "Paid / approved for future", 12), ("recipient", "Recipient", 40),
    ("recipient_city", "Recipient city", 18), ("recipient_state", "Recipient state", 8),
    ("purpose", "Purpose (as filed)", 48), ("amount", "Amount (USD)", 14),
    ("tax_year", "Tax year", 9), ("preselected_only", "Pre-selected only", 10),
    ("application_form_info", "How to apply (line 2b)", 40),
    ("application_restrictions", "Restrictions (line 2d)", 40),
    ("source_object_id", "IRS object ID", 20), ("source_url", "Filing (ProPublica)", 40),
]


def _sheet(ws, cols, rows):
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E79")
    for i, (_key, label, width) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=label)
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = width
    for r, row in enumerate(rows, 2):
        for i, (key, _label, _w) in enumerate(cols, 1):
            v = row.get(key)
            c = ws.cell(row=r, column=i, value=v)
            if isinstance(v, int):
                c.number_format = "#,##0"
            elif isinstance(v, str) and v.startswith("http"):
                c.hyperlink = v
                c.font = Font(color="0563C1", underline="single")
            c.alignment = Alignment(wrap_text=isinstance(v, str) and len(v) > 40, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(rows) + 1, 2)}"


def build(grants_csv: Path, keywords: list[str], title: str, slug: str, out_dir: Path,
          exclude_preselected: bool = False, year: int = 2026) -> dict:
    rows = read_csv(grants_csv)
    kept = list(filter_rows(rows, keywords=keywords, exclude_preselected=exclude_preselected))
    summ = summarize_by_foundation(kept)
    n_open = sum(1 for s in summ if s["preselected_only"] == "no")
    total_amt = sum(r["amount"] or 0 for r in kept)
    states = sorted({s["foundation_state"] for s in summ if s["foundation_state"]})
    built = dt.date.today().isoformat()

    out_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Foundations"
    _sheet(ws1, FOUNDATION_COLS, summ)
    ws2 = wb.create_sheet("Grants")
    _sheet(ws2, GRANT_COLS, kept)
    ws3 = wb.create_sheet("About")
    about = [
        (title, ""),
        ("Built on", built),
        ("Source", f"IRS Form 990 series e-file XML, index {year} (returns released by the IRS during {year})"),
        ("Filter", "Part XV grant rows whose recipient name or stated purpose contains: " + ", ".join(keywords)),
        ("Foundations", len(summ)), ("  of which do NOT tick 'pre-selected only'", n_open),
        ("Grant rows", len(kept)), ("Sum of matching grants (USD)", total_amt),
        ("Foundation states", ", ".join(states)),
        ("Prepared by", "an AI agent (pf-grants, open source, MIT). Not legal, tax or fundraising advice."),
        ("Check any row", "open the 'Filing (ProPublica)' link and read Part XV of that year's 990-PF."),
    ]
    for r, (k, v) in enumerate(about, 1):
        ws3.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws3.cell(row=r, column=2, value=v)
    ws3.column_dimensions["A"].width = 44
    ws3.column_dimensions["B"].width = 100

    xlsx_path = out_dir / f"{slug}.xlsx"
    wb.save(xlsx_path)
    grants_path = out_dir / f"{slug}-grants.csv"
    write_csv(kept, grants_path)
    fdn_path = out_dir / f"{slug}-foundations.csv"
    with open(fdn_path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=[k for k, _l, _w in FOUNDATION_COLS])
        w.writeheader()
        w.writerows({k: s.get(k) for k, _l, _w in FOUNDATION_COLS} for s in summ)

    readme = f"""{title}
{'=' * len(title)}

Built {built} by an AI agent with pf-grants (open source, MIT):
https://github.com/maindtim/pf-grants

WHAT THIS IS
Every US private foundation files IRS Form 990-PF. Part XV lists each grant it paid,
whether it only gives to pre-selected organizations, and how to apply if it does not.
This list is every Part XV grant row, from every 990-PF the IRS released in {year} as
e-file XML, whose recipient name or stated purpose contains: {', '.join(keywords)}.

FILES
  {slug}.xlsx               Foundations sheet (one line per foundation, sorted by matching amount),
                             Grants sheet (one line per grant), About sheet.
  {slug}-foundations.csv    Same as the Foundations sheet.
  {slug}-grants.csv         Same as the Grants sheet, all {len(GRANT_COLUMNS)} raw columns.

NUMBERS
  Foundations: {len(summ)} ({n_open} do not tick "pre-selected only")
  Grant rows:  {len(kept)}
  Sum of matching grants: USD {total_amt:,}

HOW TO USE IT
  1. Filter the Foundations sheet by your state and by "Pre-selected only" = no.
  2. Read the purpose and amounts of what they already funded; write to the application contact
     with what their own filing says they support.
  3. Before writing, open the filing link and read Part XV yourself. Recipient names and purposes
     are exactly as the foundation's preparer typed them; spelling is theirs, not ours.

HONEST LIMITS
  - Only e-filed returns released by the IRS in {year}. Paper filers are absent. Foundations that wrote
    "see attached" in Part XV appear with that text and no detail.
  - "Pre-selected only" = what the foundation ticked on line 2. Blank is treated as "no"; that is not a
    promise that unsolicited requests are welcome.
  - Keyword matching is plain substring matching on recipient and purpose. It will miss grants whose
    purpose was filed as "general support" to a literacy group whose name does not contain the word.
  - This is public data, reproduced. It is not legal, tax or fundraising advice, and it does not
    guarantee any grant.

QUESTIONS
  Open an issue at https://github.com/maindtim/pf-grants/issues . An AI agent reads and answers them.
"""
    readme_path = out_dir / f"{slug}-README.txt"
    readme_path.write_text(readme, encoding="utf-8")

    zip_path = out_dir / f"{slug}.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for p in (xlsx_path, grants_path, fdn_path):
            zf.write(p, p.name)
        zf.write(readme_path, "README.txt")

    return {"foundations": len(summ), "open": n_open, "grants": len(kept), "amount": total_amt,
            "states": len(states), "zip": str(zip_path), "xlsx": str(xlsx_path)}


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--grants", required=True)
    ap.add_argument("--keyword", action="append", required=True)
    ap.add_argument("--title", required=True)
    ap.add_argument("--slug", required=True)
    ap.add_argument("--out", default="dist")
    ap.add_argument("--exclude-preselected", action="store_true")
    ap.add_argument("--year", type=int, default=2026)
    a = ap.parse_args()
    info = build(Path(a.grants), a.keyword, a.title, a.slug, Path(a.out), a.exclude_preselected, a.year)
    print(info)
